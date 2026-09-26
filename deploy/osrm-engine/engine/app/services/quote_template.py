"""v010 正式报价单模板 —— 中越双语 · VND+RMB 双列 · 三行汇总 · 盖章位/手签位/页码。

用户 2026-09-21 拍板的三条版式决策（本文件即该决策的实现）：
  ① 对外客户版：**中越双语**，金额 **VND + RMB 双列**（并标注折算汇率）；
  ② 明细颗粒度：客户版只给 **运输费 / 口岸费 / 总计** 三行（不逐项拆给客户）；
  ③ 版面含 **公司盖章位 + 客户手签位 + 页脚页码（第 X 页 / 共 Y 页）**。

口径（与既有 quote_docx_builder 一致，勿另立）：
  * 客户版 `payload.price_vnd` = 运输服务费（门到门，**已含利润**）；
  * 口岸费 = `payload.border_fees.total_vnd`（或 total_rmb × 汇率）；
  * 总计 = 运输费 + 口岸费；
  * **客户版绝不出现「成本 / 利润 / 毛利」字样与内部备注**，内部版才有。

本模块是**模板草案**：先用它出样张给用户确认，确认后再让
`quote_docx_builder.build()` 委托到本模块（届时 `POST /quote/export` 自动升级，接口不变）。
"""
from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.services._paths import data_dir, resource_dir
# 复用共用的金额格式化与中文大写（v010 起两版式共用 quote_format，避免循环导入）
from app.services.quote_format import _fmt_vnd, _vnd_uppercase

DEFAULT_RATE = 3891.05          # VND per 1 RMB —— 取不到实时汇率时的兜底（与 data/exchange_rate.json 一致）

# ══════════════════════ 基础排版工具 ══════════════════════


def _set_font(run, size=10.5, bold=False, cn="宋体"):
    """西文用 Times New Roman（越南语带声调符号必须走它），中文 eastAsia 用宋体。"""
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), cn)


def _bi(cn: str, vi: str) -> str:
    """双语标签：中文 / Tiếng Việt。"""
    return f"{cn} / {vi}"


def _para(doc, text="", *, size=10.5, bold=False, align=None, cn="宋体", space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        _set_font(p.add_run(text), size=size, bold=bold, cn=cn)
    return p


def _cell_text(cell, text, *, size=10, bold=False, align=None, cn="宋体"):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    _set_font(p.add_run(str(text)), size=size, bold=bold, cn=cn)
    return cell


def _shade(cell, color="EDEDED"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _field(run, instr: str) -> None:
    """插入 Word 域（用于页码/总页数）。"""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr_el, end):
        run._r.append(el)


# ══════════════════════ 越南语数字大写（Bằng chữ） ══════════════════════

_VI_UNITS = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]


def _vi_three(n: int, full: bool) -> str:
    """0..999 → chữ。full=True 表示更高位已出现（需补 'không trăm' / 'lẻ'）。"""
    tr, ch, dv = n // 100, (n // 10) % 10, n % 10
    parts: list[str] = []
    if tr or full:
        parts.append(_VI_UNITS[tr] + " trăm")
    if ch == 0:
        if dv and (tr or full):
            parts.append("lẻ")
        if dv:
            parts.append(_VI_UNITS[dv])
    elif ch == 1:
        parts.append("mười")
        if dv == 5:
            parts.append("lăm")
        elif dv:
            parts.append(_VI_UNITS[dv])
    else:
        parts.append(_VI_UNITS[ch] + " mươi")
        if dv == 1:
            parts.append("mốt")
        elif dv == 5:
            parts.append("lăm")
        elif dv == 4:
            parts.append("tư")
        elif dv:
            parts.append(_VI_UNITS[dv])
    return " ".join(parts)


def vnd_words_vi(amount) -> str:
    """越南语金额大写：12345678 → 'Mười hai triệu ba trăm bốn mươi lăm nghìn sáu trăm bảy mươi tám đồng'。

    ⚠️ 样张阶段请越南同事核一遍用词（mươi/lăm/mốt/tư 的边界值最易错）。
    """
    n = int(round(float(amount or 0)))
    if n == 0:
        return "Không đồng"
    scales = ["", "nghìn", "triệu", "tỷ"]
    raw: list[tuple[int, str]] = []
    i = 0
    while n > 0:
        raw.append((n % 1000, scales[i] if i < len(scales) else "tỷ"))
        n //= 1000
        i += 1
    # 从最高位往低渲染：idx>0 表示"前面已有更高位" → 该组需要补 "không trăm" / "lẻ"
    # （踩过的坑：若按低位累加时的 any(groups) 判断，最高位也会被补成 "không trăm bảy mươi hai"）
    parts = []
    for idx, (g, scale) in enumerate(reversed(raw)):
        if not g:
            continue
        parts.append(f"{_vi_three(g, full=(idx > 0))} {scale}".strip())
    txt = " ".join(parts)
    return txt[0].upper() + txt[1:] + " đồng"


# ══════════════════════ 金额 / 汇率 ══════════════════════


def _rate_of(payload, rate_vnd_per_rmb: float | None) -> float:
    if rate_vnd_per_rmb and rate_vnd_per_rmb > 0:
        return float(rate_vnd_per_rmb)
    bf = getattr(payload, "border_fees", None)
    r = getattr(bf, "exchange_rate", None) if bf else None
    try:
        r = float(r) if r else 0.0
    except (TypeError, ValueError):
        r = 0.0
    return r if r > 0 else DEFAULT_RATE


def _rmb(vnd: float, rate: float) -> float:
    return float(vnd or 0) / rate


def fee_totals(payload, rate: float) -> dict:
    """客户版三行口径 + 内部版补充口径，全部 VND。"""
    transport = float(getattr(payload, "price_vnd", 0) or 0)
    bf = getattr(payload, "border_fees", None)
    border_vnd = 0.0
    if bf is not None:
        border_vnd = float(getattr(bf, "total_vnd", 0) or 0)
        if not border_vnd:
            border_vnd = float(getattr(bf, "total_rmb", 0) or 0) * rate
    return {
        "transport_vnd": transport,
        "border_vnd": border_vnd,
        "total_vnd": transport + border_vnd,
        "cost_vnd": float(getattr(payload, "cost_total", 0) or 0),
        "profit_vnd": float(getattr(payload, "profit_vnd", 0) or 0),
    }


# ══════════════════════ 各区块 ══════════════════════


def _add_letterhead(doc, profile) -> None:
    _para(doc, profile.get("cn_name", ""), size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, cn="黑体", space_after=0)
    _para(doc, profile.get("vi_name", ""), size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _para(doc, profile.get("address_cn", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _para(doc, profile.get("address_vi", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _para(doc, "电话/Tel: {p}    邮箱/Email: {e}    税号/MST: {t}".format(
        p=profile.get("phone", ""), e=profile.get("email", ""), t=profile.get("tax_id", "")),
        size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def _add_title(doc, quote_no: str, date_str: str) -> None:
    _para(doc, "运输服务报价单", size=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, cn="黑体", space_after=0)
    _para(doc, "BẢNG BÁO GIÁ DỊCH VỤ VẬN TẢI", size=12, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    _para(doc, _bi(f"报价单号：{quote_no}", f"Số báo giá: {quote_no}"), size=10,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _para(doc, _bi(f"报价日期：{date_str}", f"Ngày báo giá: {date_str}"), size=10,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def _add_basic_info(doc, payload, date_str: str) -> None:
    rows = [
        (_bi("客户名称", "Khách hàng"), payload.customer_name),
        (_bi("运输线路", "Tuyến đường"), payload.route_name),
        (_bi("参考里程", "Quãng đường tham khảo"),
         (f"{payload.distance_km:,.1f} km" if payload.distance_km else "—")
         + (f" ｜ {_bi('预计时效', 'Thời gian dự kiến')}: {payload.duration_h:.1f} h" if payload.duration_h else "")),
        (_bi("车型 / 车数", "Loại xe / Số lượng"), f"{payload.vehicle_model_name} × {payload.vehicle_count}"),
        (_bi("报价有效期", "Hiệu lực báo giá"), _bi(f"{payload.validity_days} 天", f"{payload.validity_days} ngày")),
    ]
    tab = doc.add_table(rows=len(rows), cols=2)
    tab.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        _shade(tab.cell(i, 0), "F5F5F5")
        _cell_text(tab.cell(i, 0), k, size=10, bold=True)
        _cell_text(tab.cell(i, 1), v, size=10)
    tab.columns[0].width = Cm(5.2)
    tab.columns[1].width = Cm(11.3)
    _para(doc, "", space_after=2)


def _add_cargo(doc, payload) -> None:
    if not payload.cargo:
        return
    _para(doc, _bi("货物明细", "Chi tiết hàng hóa"), size=11, bold=True, cn="黑体")
    tab = doc.add_table(rows=1 + len(payload.cargo), cols=5)
    tab.style = "Table Grid"
    heads = [_bi("货物名称", "Tên hàng"), _bi("件数", "Số kiện"), _bi("单件重量(吨)", "TL/kiện (tấn)"),
             _bi("尺寸(长×宽×高 cm)", "Kích thước (D×R×C cm)"), _bi("合计重量(吨)", "Tổng TL (tấn)")]
    for j, h in enumerate(heads):
        _shade(tab.cell(0, j), "EDEDED")
        _cell_text(tab.cell(0, j), h, size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, c in enumerate(payload.cargo, start=1):
        size = "—" if not (c.l_cm and c.w_cm and c.h_cm) else f"{c.l_cm:g}×{c.w_cm:g}×{c.h_cm:g}"
        _cell_text(tab.cell(i, 0), c.name, size=9.5)
        _cell_text(tab.cell(i, 1), f"{c.qty:g}", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(tab.cell(i, 2), f"{c.weight_ton:g}", size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
        _cell_text(tab.cell(i, 3), size, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(tab.cell(i, 4), f"{c.weight_ton * c.qty:g}", size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(doc, "", space_after=2)


def _add_fee_summary(doc, payload, rate: float, *, role: str) -> None:
    """费用汇总：客户版三行（运输费 / 口岸费 / 总计），VND + RMB 双列。

    大件/超限单（`manual_oversize_fee_vnd > 0`）：运输行文案沿用既有对客户说法
    「大件/超限运输专线价（手工录入）」，并在表下打印 `oversize_note` 说明
    —— 既有测试 `tests/test_oversize.py::test_docx_shows_manual_oversize_trip_fee` 守着这条口径。
    """
    t = fee_totals(payload, rate)
    manual = float(getattr(payload, "manual_oversize_fee_vnd", 0) or 0)
    _para(doc, _bi("费用明细", "Chi tiết chi phí"), size=11, bold=True, cn="黑体")
    tab = doc.add_table(rows=2 + 3, cols=3)
    tab.style = "Table Grid"
    heads = [_bi("费用项目", "Khoản mục"), _bi("金额（VND）", "Số tiền (VND)"),
             _bi("金额（RMB）", "Số tiền (RMB)")]
    for j, h in enumerate(heads):
        _shade(tab.cell(0, j), "EDEDED")
        _cell_text(tab.cell(0, j), h, size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    rows = [
        (_bi("运输服务费（门到门）" if not manual else "大件/超限运输专线价（手工录入）",
             "Phí vận tải (door-to-door)" if not manual else "Phí vận tải tuyến hàng quá khổ (nhập tay)"),
         t["transport_vnd"], False),
        (_bi("进出口口岸费", "Phí cửa khẩu xuất nhập khẩu"), t["border_vnd"], False),
        (_bi("总计", "Tổng cộng"), t["total_vnd"], True),
    ]
    for i, (name, vnd, bold) in enumerate(rows, start=1):
        _cell_text(tab.cell(i, 0), name, size=10, bold=bold)
        _cell_text(tab.cell(i, 1), _fmt_vnd(vnd), size=10, bold=bold, align=WD_ALIGN_PARAGRAPH.RIGHT)
        _cell_text(tab.cell(i, 2), _fmt_vnd(_rmb(vnd, rate)), size=10, bold=bold,
                   align=WD_ALIGN_PARAGRAPH.RIGHT)
        if bold:
            for j in range(3):
                _shade(tab.cell(i, j), "F5F5F5")
    _para(doc, _bi(f"折合汇率：1 RMB = {rate:,.2f} VND（汇率波动以实际结算日为准）",
                   f"Tỷ giá quy đổi: 1 RMB = {rate:,.2f} VND (tỷ giá có thể thay đổi theo ngày thanh toán)"),
          size=9, space_after=2)
    _para(doc, _bi(f"总计大写（人民币）：{_vnd_uppercase(_rmb(t['total_vnd'], rate))}",
                   f"Tổng cộng bằng chữ (VND): {vnd_words_vi(t['total_vnd'])}"), size=9.5, bold=True, space_after=8)
    # 大件/超限：把手写说明原样带上（对客户解释为什么不是常规价）
    if manual > 0 and str(getattr(payload, "oversize_note", "") or "").strip():
        _para(doc, str(payload.oversize_note).strip(), size=9, space_after=8)

    if role == "internal":
        _para(doc, _bi("【内部版】成本与利润", "[Nội bộ] Chi phí và lợi nhuận"), size=11, bold=True, cn="黑体")
        it = doc.add_table(rows=5, cols=2)
        it.style = "Table Grid"
        bd = payload.breakdown or {}
        gross = (t["profit_vnd"] / t["total_vnd"] * 100) if t["total_vnd"] else 0
        irows = [
            (_bi("距离成本", "Chi phí quãng đường"), bd.get("cost_distance", 0)),
            (_bi("油耗成本", "Chi phí nhiên liệu"), bd.get("cost_fuel", 0)),
            (_bi("路桥费", "Phí cầu đường"), bd.get("cost_toll", 0)),
            (_bi("装卸费", "Phí bốc xếp"), bd.get("cost_loading", 0)),
            (_bi("成本合计 / 利润 / 毛利率", "Tổng chi phí / Lợi nhuận / Biên LN"),
             f"{_fmt_vnd(t['cost_vnd'])} / {_fmt_vnd(t['profit_vnd'])} / {gross:.1f}%"),
        ]
        for i, (k, v) in enumerate(irows):
            _cell_text(it.cell(i, 0), k, size=9.5, bold=True)
            _cell_text(it.cell(i, 1), v if isinstance(v, str) else _fmt_vnd(v), size=9.5,
                       align=WD_ALIGN_PARAGRAPH.RIGHT)
        _para(doc, "", space_after=2)


def _add_payment_block(doc, profile) -> None:
    """收款账户（中越双语）+（可选的）付款条件。

    ⚠️ 付款条件的口径（用户 2026-09-21 指示）：**具体付款条款在正式合同中约定**，
       报价单只写一句"付款方式在正式合同中约定"；若 `payment_terms_cn/vi` 配空，
       则整个「付款条件」小标题都不出现（只留收款账户）。
    ⚠️ 措辞注意：报银行/海关/对外合同**禁用「保证金」**（既有铁律），本区块只写预付款/尾款。
    """
    terms_cn = str(profile.get("payment_terms_cn", "") or "").strip()
    terms_vi = str(profile.get("payment_terms_vi", "") or "").strip()
    if terms_cn or terms_vi:
        _para(doc, _bi("付款条件", "Điều khoản thanh toán"), size=11, bold=True, cn="黑体")
        _para(doc, _bi(terms_cn, terms_vi), size=9.5, space_after=6)

    _para(doc, _bi("收款账户", "Tài khoản nhận tiền"), size=11, bold=True, cn="黑体")
    rows = [
        (_bi("收款单位", "Đơn vị thụ hưởng"),
         _bi(profile.get("bank_holder_cn", ""), profile.get("bank_holder_vi", ""))),
        (_bi("开户银行", "Ngân hàng"),
         _bi(profile.get("bank_name_cn", ""), profile.get("bank_name_vi", ""))),
        (_bi("越南盾账户", "Tài khoản VND"), profile.get("bank_account_vnd", "")),
        (_bi("美金账户", "Tài khoản USD"), profile.get("bank_account_usd", "")),
        (_bi("SWIFT Code", "Mã SWIFT"), profile.get("bank_swift", "")),
    ]
    rows = [(k, v) for k, v in rows if str(v).strip()]
    tab = doc.add_table(rows=len(rows), cols=2)
    tab.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        _shade(tab.cell(i, 0), "F5F5F5")
        _cell_text(tab.cell(i, 0), k, size=9.5, bold=True)
        _cell_text(tab.cell(i, 1), v, size=9.5)
    tab.columns[0].width = Cm(4.6)
    tab.columns[1].width = Cm(11.9)
    _para(doc, _bi("转账时请在备注中注明报价单号，以便核对到账。",
                   "Khi chuyển khoản vui lòng ghi rõ số báo giá để đối chiếu."), size=9, space_after=6)


def _add_terms(doc, payload) -> None:
    _para(doc, _bi("其他约定", "Các thỏa thuận khác"), size=11, bold=True, cn="黑体")
    terms = [
        _bi("1. 本报价为门到门全程运输，含上述运输费与进出口口岸费，不含关税、增值税及海关查验产生的仓储费用。",
            "1. Báo giá theo hình thức vận tải door-to-door, đã gồm phí vận tải và phí cửa khẩu nêu trên; "
            "chưa gồm thuế nhập khẩu, VAT và phí lưu kho phát sinh khi kiểm hóa."),
        _bi("2. 报价有效期见上表；超期需重新确认运价。",
            "2. Hiệu lực báo giá theo bảng trên; quá hạn vui lòng xác nhận lại giá."),
        _bi("3. 货物重量/尺寸以装车时实际测量为准，若与申报不符，运价按实际重算。",
            "3. Khối lượng/kích thước hàng hóa tính theo thực tế khi xếp xe; nếu khác khai báo, "
            "giá sẽ được tính lại theo thực tế."),
    ]
    for t in terms:
        _para(doc, t, size=9.5, space_after=2)
    _para(doc, "", space_after=6)


def _asset_candidates(fname: str) -> list[Path]:
    """图片资产的搜索路径（按优先级）。

    踩过的坑：开发态 `resource_dir()` 指向 `backend/`（不是仓库根 `resources/`），
    只有打包态才指向 `%APPDATA%/aiosrm-desktop/resources`；所以必须多路径兜底，
    否则"把图放到 resources/ 就会自动生效"在开发态根本不成立。
    """
    here = Path(__file__).resolve()
    repo = here.parents[3]            # …/AIOSRM++/
    out = [
        data_dir() / fname,                  # 用户数据目录（可自行替换）
        resource_dir() / fname,              # 打包态 exe 同级 resources/ 或 OSRM_RESOURCE_DIR
        here.parents[2] / fname,             # 开发态 backend/
        repo / "resources" / fname,          # 开发态 仓库根 resources/ ← 放这里
        repo / fname,
    ]
    if getattr(sys, "frozen", False):
        # ⚠️ 打包态兜底（v010 真机验收抓到的坑）：冻结后 __file__ 在 PyInstaller 包里，
        # 上面几条全指不到 extraResources 落地处 → 章盖不出来（显示【此处盖章】）。
        # exe = <app>/resources/backend/AIOSRM-backend.exe → parents[1] = <app>/resources/
        exe = Path(sys.executable).resolve()
        out[2:2] = [exe.parents[1] / "resources" / fname, exe.parents[1] / fname]
    return out


def _asset_path(fname: str) -> Path | None:
    for p in _asset_candidates(fname):
        try:
            if p.exists():
                return p
        except OSError:
            continue
    return None


def _seal_candidates() -> list[Path]:
    """公司章 png 的搜索路径（保留旧名：测试与门禁引用它）。"""
    return _asset_candidates("company_seal.png")


def _seal_path() -> Path | None:
    return _asset_path("company_seal.png")


def _sign_candidates() -> list[Path]:
    """手写签名 png 的搜索路径（v010：放到 resources/company_sign.png 即自动进手签位）。"""
    return _asset_candidates("company_sign.png")


def _sign_path() -> Path | None:
    return _asset_path("company_sign.png")


def _sign_enabled(profile: dict) -> bool:
    """手写签名开关（读 `company_info.json` 的 `signature_enabled`）。

    用户 2026-09-21：「就只保留公司签章，后面再改」→ 默认 false（只盖公章）。
    要恢复手签：把配置项改成 true（并确保 `resources/company_sign.png` 在位），不用改代码。
    """
    v = profile.get("signature_enabled", False)
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in {"1", "true", "yes", "on", "y", "是"}


def _add_sign_block(doc, profile, quote_no: str) -> None:
    """盖章位 + 客户手签位（替代纯文字落款，便于盖章与回传）。"""
    _para(doc, _bi("签署与确认", "Ký kết và xác nhận"), size=11, bold=True, cn="黑体")
    tab = doc.add_table(rows=3, cols=2)
    tab.style = "Table Grid"
    _shade(tab.cell(0, 0), "F5F5F5")
    _shade(tab.cell(0, 1), "F5F5F5")
    _has_sign = _sign_enabled(profile) and _sign_path() is not None
    _cell_text(tab.cell(0, 0),
               _bi(f"{profile.get('cn_name', '')}（签字并盖章）" if _has_sign
                   else f"{profile.get('cn_name', '')}（加盖公章）",
                   f"{profile.get('vi_name', '')} (ký tên và đóng dấu)" if _has_sign
                   else f"{profile.get('vi_name', '')} (đóng dấu)"),
               size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell_text(tab.cell(0, 1), _bi("客户确认（签字/盖章）", "Xác nhận của khách hàng (ký, đóng dấu)"),
               size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    seal = _seal_path()
    sign = _sign_path() if _has_sign else None     # 关掉手签时只上公章（开关见 _sign_enabled）
    c = tab.cell(1, 0)
    c.text = ""
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 报价方落款：手签（上）与公章（下）**分两行**居中 —— 同段内联会并排，看着挤也不是惯例
    if sign is not None:
        try:
            p.add_run().add_picture(str(sign), width=Cm(4.0))
        except Exception:  # noqa: BLE001 —— 图坏了也要能出单
            pass
    p_seal = c.add_paragraph() if sign is not None else p
    p_seal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if seal is not None:
        try:
            p_seal.add_run().add_picture(str(seal), width=Cm(3.6))
        except Exception:  # noqa: BLE001 —— 图片坏了也要能出单
            _set_font(p_seal.add_run("【此处盖章 / Đóng dấu】"), size=10)
    else:
        # 占位符只写「此处盖章」，不要在客户可见文档里写内部说明（章图放置路径见 _asset_candidates 注释）
        _set_font(p_seal.add_run("【此处盖章 / Đóng dấu】"), size=10)
    for _ in range(2):
        _set_font(c.add_paragraph().add_run("\n"), size=10)

    c2 = tab.cell(1, 1)
    c2.text = ""
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(4):
        _set_font(p2.add_run("\n"), size=10)
    _set_font(c2.add_paragraph().add_run(_bi("签字 / Chữ ký：______________", "Họ tên: ______________")), size=9)

    _cell_text(tab.cell(2, 0), _bi("日期 / Ngày：____/____/______", f"Số báo giá: {quote_no}"), size=9)
    _cell_text(tab.cell(2, 1), _bi("日期 / Ngày：____/____/______", "(客户回传后生效 / có hiệu lực sau khi khách xác nhận)"), size=9)


def _add_footer(doc, profile, quote_no: str) -> None:
    """页脚：公司名 ｜ 单号 ｜ 第 X 页 / 共 Y 页（Trang X/Y）—— 页码用 Word 域，自动更新。"""
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run(f"{profile.get('cn_name', '')} ｜ {quote_no} ｜ "), size=8)
    _set_font(p.add_run("第 "), size=8)
    _field(p.add_run(), "PAGE")
    _set_font(p.add_run(" 页 / 共 "), size=8)
    _field(p.add_run(), "NUMPAGES")
    _set_font(p.add_run(" 页（Trang "), size=8)
    _field(p.add_run(), "PAGE")
    _set_font(p.add_run("/"), size=8)
    _field(p.add_run(), "NUMPAGES")
    _set_font(p.add_run("）"), size=8)


# ══════════════════════ 入口 ══════════════════════


def build(payload, profile: dict, quote_no: str, *, rate_vnd_per_rmb: float | None = None,
          date_str: str | None = None) -> bytes:
    """生成 v010 版式报价单（role=customer 双语三行汇总；role=internal 额外给成本利润）。"""
    import datetime

    date_str = date_str or datetime.date.today().strftime("%d/%m/%Y")
    rate = _rate_of(payload, rate_vnd_per_rmb)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.6)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    _add_letterhead(doc, profile)
    _add_title(doc, quote_no, date_str)
    _add_basic_info(doc, payload, date_str)
    _add_cargo(doc, payload)
    _add_fee_summary(doc, payload, rate, role=payload.role)
    _add_payment_block(doc, profile)
    _add_terms(doc, payload)
    _add_sign_block(doc, profile, quote_no)
    _add_footer(doc, profile, quote_no)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
