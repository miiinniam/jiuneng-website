"""v013 装箱单 docx —— 中越双语 · 逐车一页 · 该车摆放清单 + 3D 装车示意图。

用户 2026-09-23 拍板（见 `docs/版本档案/v013-装车系统升级-2026-09-23/实施计划.md` T4）：
  ① 由 v007 的 CSV 升级为 **docx 可打印版式**；
  ② **逐车一页**：车型 / 车厢尺寸 / 限重 / 总重 / 装载率 + 该车摆放清单表 + 该车 3D 示意图；
  ③ 默认**中越双语**（D5），与报价单同一套版式工具；
  ④ 客户版（`role=client`）正文**绝不出现「成本 / 利润 / 毛利」**字样 —— 沿用报价单既有铁律。

单位与坐标口径（**与前端 3D / `packer.py` 完全一致，勿另立**）：
  * `x/y/z` = 车厢坐标 **cm**，原点 = **前壁左下角**；`x` 沿车长（前壁→车门）、`y` = 高度、`z` = 宽度；
  * **`y == 0` 表示落地** → 层列写「地板 / Đáy」；其余按**高度自下而上的次序**写「第N层 / Tầng N」
    （N = 该车所有 y>0 的不同高度值升序去重后的序号，不是 `y/层高` 的猜测值）；
  * `dx/dy/dz` = 该件沿 x/y/z 的占用尺寸；表格「尺寸」列按 **长(dx)×宽(dz)×高(dy)** 排
    （人对箱子的读法），并在表下明写坐标系，避免 x/y/z 顺序被误读。

踩过的坑（都在 `tests/test_packing_list.py` 里锁住）：
  * **同字节的图会被 python-docx 去重成一个 `word/media/*` 部件** → 「带图车辆数 == media 张数」
    的用例必须给**不同字节**的示意图，否则断言假失败；
  * `snapshot_png` 允许 `data:image/png;base64,…` 前缀，也允许**纯 base64**（两种都支持）；
  * 无示意图的车**不插图、不留标题、不留空页**（缺省即整段省略）；
  * `trucks == []` → 由 `build()` 抛 `ValueError`，端点折成 **400**（不是空单页，见 `api/packing.py`）。
"""
from __future__ import annotations

import base64
import binascii
import datetime
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# 复用报价单的版式工具（模块内私有名，**刻意复用**：两种单据的字体/信头/页脚/表格样式
# 必须同源，各写一套迟早漂移；报价单侧的版式改动会自动传导到装箱单）。
from app.services.quote_template import (
    _add_footer,
    _add_letterhead,
    _bi,
    _cell_text,
    _para,
    _set_font,
    _shade,
)

# 端点把它折成 400（见 `app/api/packing.py`）
EMPTY_TRUCKS_DETAIL = "装箱单生成失败：没有车辆数据（trucks 为空）"

_TITLE_CN = "装箱单"
_TITLE_VI = "BẢNG KÊ ĐÓNG GÓI (PACKING LIST)"

# 车型英文/缺省名 → 中越展示名（`truck.type` 缺省或未知 → 货车 / Xe tải）
_TRUCK_TYPE_BI = {
    "flatbed": ("平板车", "Xe sàn"),
    "lowbed": ("低平板车", "Xe sàn thấp"),
    "box": ("厢式车", "Xe thùng"),
    "container": ("集装箱车", "Xe container"),
}

# 车厢坐标系说明（每车页表下打印一次：x/y/z 的顺序最容易被误读成 长/宽/高）
_AXIS_NOTE_BI = _bi(
    "坐标系：x 沿车长（前壁→车门）/ y 高度 / z 宽度，原点 = 前壁左下角，单位 cm；"
    "「尺寸」列按 长(dx)×宽(dz)×高(dy) 排列；层列 y=0 为地板，其余按高度自下而上次序编号。",
    "Hệ tọa độ: x = chiều dài thùng xe (vách trước → cửa sau) / y = chiều cao / z = chiều rộng, "
    "gốc tọa độ = góc dưới bên trái vách trước, đơn vị cm; cột \"Kích thước\" xếp theo "
    "Dài(dx)×Rộng(dz)×Cao(dy); ở cột Tầng, y=0 là sàn xe, các tầng khác đánh số từ thấp lên cao.",
)

# 图片最大宽度（A4 正文净宽 = 21.0 - 2.0×2 = 17.0 cm，留些余量免贴近页边）
_IMG_WIDTH_CM = 15.5

_ILLEGAL_FN = re.compile(r'[\\/:*?"<>|\r\n\t]+')


# ══════════════════════ 取值 / 归一化（dict 与 pydantic 模型都能吃） ══════════════════════


def _g(obj, key, default=None):
    """取字段：dict 与 pydantic 模型通用（缺省字段全给默认值，前端少传不炸）。"""
    if obj is None:
        return default
    v = obj.get(key, default) if isinstance(obj, dict) else getattr(obj, key, default)
    return default if v is None else v


def _f(v, default=0.0) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _num(v) -> str:
    """尺寸/坐标显示：整数不带小数点，非整数原样（12.5 → 12.5）。"""
    f = _f(v)
    return f"{f:.0f}" if abs(f - round(f)) < 1e-9 else f"{f:g}"


def _kg(v) -> str:
    f = _f(v)
    return f"{f:,.0f}" if abs(f - round(f)) < 1e-9 else f"{f:,.1f}"


def decode_png(raw) -> bytes | None:
    """解析示意图：`data:image/png;base64,…` 前缀 / 纯 base64 都支持；解不出来 → None。

    刻意**不抛异常**：图坏了也要能出单（与报价单盖章图同一条纪律）。
    """
    s = str(raw or "").strip()
    if not s:
        return None
    if s.lower().startswith("data:"):
        if "," not in s:                      # data:image/png;base64（没有负载）
            return None
        s = s.split(",", 1)[1]
    s = re.sub(r"\s+", "", s)
    if not s or not re.fullmatch(r"[A-Za-z0-9+/]*={0,2}", s):
        return None
    s += "=" * (-len(s) % 4)                  # 容错：前端手工截断过的串也能补回填充
    try:
        blob = base64.b64decode(s, validate=True)
    except (binascii.Error, ValueError):
        return None
    return blob or None


def _norm_placement(p) -> dict:
    return {
        "cargoId": str(_g(p, "cargoId", "") or ""),
        "name": str(_g(p, "name", "") or "").strip() or _bi("货物", "Hàng hóa"),
        "x": _f(_g(p, "x", 0)), "y": _f(_g(p, "y", 0)), "z": _f(_g(p, "z", 0)),
        "dx": _f(_g(p, "dx", 0)), "dy": _f(_g(p, "dy", 0)), "dz": _f(_g(p, "dz", 0)),
        "weight": _f(_g(p, "weight", 0)),
        "stackable": bool(_g(p, "stackable", True)),   # 缺省 True（旧数据零影响）
    }


def _norm_unplaced(u) -> dict:
    """未装配件（`reason` 缺省就只写货名，不臆造原因）。"""
    name = str(_g(u, "name", "") or "").strip() or _bi("货物", "Hàng hóa")
    return {
        "cargoId": str(_g(u, "cargoId", "") or ""),
        "name": name,
        "reason": str(_g(u, "reason", "") or "").strip(),
    }


def _norm_truck(t, fallback_index: int = 1) -> dict:
    idx = _g(t, "index", None)
    try:
        index = int(idx) if idx is not None else fallback_index
    except (TypeError, ValueError):
        index = fallback_index
    tr = _g(t, "truck", {}) or {}
    truck = {k: _f(_g(tr, k, 0)) for k in ("L", "W", "H", "maxWeight")}
    truck["type"] = str(_g(tr, "type", "") or "")
    placements = [_norm_placement(p) for p in (_g(t, "placements", []) or [])]
    unplaced = [_norm_unplaced(u) for u in (_g(t, "unplaced", []) or [])]

    rate = _g(t, "loading_rate", None)
    rate = None if rate is None else _f(rate, -1.0)
    if rate is not None and rate >= 0:
        if rate > 1.0:                        # 容错：传百分比（77.3）而不是小数（0.773）
            rate = rate / 100.0 if rate <= 100.0 else -1.0
        rate = min(max(rate, 0.0), 1.0) if rate >= 0 else None
    else:
        rate = None

    return {
        "index": index,
        "model_name": str(_g(t, "model_name", "") or "").strip(),
        "loading_rate": rate,
        "truck": truck,
        "placements": placements,
        "unplaced": unplaced,
        "png": decode_png(_g(t, "snapshot_png", None)),
    }


# ══════════════════════ 计算口径 ══════════════════════


def truck_volume_cm3(t: dict) -> float:
    tr = t["truck"]
    return tr["L"] * tr["W"] * tr["H"]


def used_volume_cm3(t: dict) -> float:
    return sum(p["dx"] * p["dy"] * p["dz"] for p in t["placements"])


def total_weight_kg(t: dict) -> float:
    return sum(p["weight"] for p in t["placements"])


def loading_rate(t: dict) -> float | None:
    """装载率：显式传值优先；缺省按 **摆放件体积 / 车厢容积** 算（无容积 → None）。"""
    if t["loading_rate"] is not None:
        return t["loading_rate"]
    vol = truck_volume_cm3(t)
    return (used_volume_cm3(t) / vol) if vol > 0 else None


def model_label(t: dict) -> str:
    """车型名：显式 `model_name` 优先；缺省用车型 + 车厢尺寸拼（用户可读，不留空）。"""
    if t["model_name"]:
        return t["model_name"]
    tr = t["truck"]
    cn, vi = _TRUCK_TYPE_BI.get(tr["type"], ("货车", "Xe tải"))
    if tr["L"] > 0 and tr["W"] > 0 and tr["H"] > 0:
        dims = f"{tr['L'] / 100:.1f}m×{tr['W'] / 100:.1f}m×{tr['H'] / 100:.1f}m"
        return _bi(f"{cn} {dims}", f"{vi} {dims}")
    return _bi(cn, vi)


def layer_labels(t: dict) -> dict:
    """y>0 的不同高度值升序去重 → {高度: 层号}（y=0 恒为「地板」，不进这张表）。"""
    ys = sorted({round(p["y"], 3) for p in t["placements"] if p["y"] > 0.01})
    return {y: i + 1 for i, y in enumerate(ys)}


def layer_text(t: dict, p: dict, levels: dict) -> str:
    if p["y"] <= 0.01:
        return _bi("地板", "Đáy")
    n = levels.get(round(p["y"], 3))
    if n is None:                              # 理论上不会发生（levels 由同一批数据生成）
        n = 1
    return _bi(f"第{n}层", f"Tầng {n}")


def _truck_summary_line(t: dict) -> str:
    total = total_weight_kg(t)
    rate = loading_rate(t)
    return _bi(
        f"共 {len(t['placements'])} 件，总重 {_kg(total)} kg"
        + (f"，装载率 {rate * 100:.1f}%" if rate is not None else ""),
        f"Tổng {len(t['placements'])} kiện, tổng trọng lượng {_kg(total)} kg"
        + (f", tỷ lệ xếp hàng {rate * 100:.1f}%" if rate is not None else ""),
    )


def filename_for(payload) -> str:
    """`装箱单_起点_终点_日期.docx`（非法文件名字符去掉，缺项写「未填写」）。"""
    job = _g(payload, "job", {}) or {}

    def seg(v, default):
        s = _ILLEGAL_FN.sub("", str(v or "")).strip()
        return s or default

    origin = seg(_g(job, "origin", ""), "未填写")
    dest = seg(_g(job, "destination", ""), "未填写")
    date = seg(_g(job, "date", ""), datetime.date.today().strftime("%Y-%m-%d"))
    return f"装箱单_{origin}_{dest}_{date}.docx"


# ══════════════════════ 版面 ══════════════════════


def _page_setup(doc) -> None:
    """A4 + 页边距 + Normal 字体（与报价单同一套数值）。"""
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.6)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _add_title_block(doc, *, is_client: bool, date_str: str) -> None:
    _para(doc, _TITLE_CN, size=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, cn="黑体", space_after=0)
    _para(doc, _TITLE_VI, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    _para(doc, _bi(f"装箱日期：{date_str}", f"Ngày đóng hàng: {date_str}"), size=10,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6 if not is_client else 10)
    if not is_client:
        # 内部版才有的一句；客户版绝不出现（连同「成本/利润/毛利」一起被测试锁住）
        _para(doc, _bi("【内部版】仅供内部核算使用，勿对外发送",
                       "[Bản nội bộ] Chỉ dùng cho hạch toán nội bộ, không gửi ra bên ngoài."),
              size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def _add_job_info(doc, job, trucks) -> None:
    origin = str(_g(job, "origin", "") or "").strip() or "—"
    dest = str(_g(job, "destination", "") or "").strip() or "—"
    border = str(_g(job, "border_crossing", "") or "").strip() or "—"
    date_str = str(_g(job, "date", "") or "").strip() or "—"
    total_w = sum(total_weight_kg(t) for t in trucks)
    total_pcs = sum(len(t["placements"]) for t in trucks)
    rows = [
        (_bi("起点", "Khởi điểm"), origin),
        (_bi("终点", "Điểm đến"), dest),
        (_bi("口岸", "Cửa khẩu"), border),
        (_bi("日期", "Ngày"), date_str),
        (_bi("车数 / 件数 / 总重", "Số xe / Số kiện / Tổng trọng lượng"),
         _bi(f"{len(trucks)} 车 / {total_pcs} 件 / {_kg(total_w)} kg",
             f"{len(trucks)} xe / {total_pcs} kiện / {_kg(total_w)} kg")),
    ]
    tab = doc.add_table(rows=len(rows), cols=2)
    tab.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        _shade(tab.cell(i, 0), "F5F5F5")
        _cell_text(tab.cell(i, 0), k, size=10, bold=True)
        _cell_text(tab.cell(i, 1), v, size=10)
    tab.columns[0].width = Cm(5.6)
    tab.columns[1].width = Cm(11.4)
    _para(doc, "", space_after=2)


def _add_summary_table(doc, trucks) -> None:
    _para(doc, _bi("车辆汇总", "Tổng hợp xe"), size=11, bold=True, cn="黑体")
    heads = [_bi("车号", "Số xe"), _bi("车型", "Loại xe"), _bi("件数", "Số kiện"),
             _bi("总重(kg)", "Tổng TL (kg)"), _bi("装载率", "Tỷ lệ xếp hàng")]
    tab = doc.add_table(rows=1 + len(trucks) + 1, cols=len(heads))
    tab.style = "Table Grid"
    for j, h in enumerate(heads):
        _shade(tab.cell(0, j), "EDEDED")
        _cell_text(tab.cell(0, j), h, size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, t in enumerate(trucks, start=1):
        rate = loading_rate(t)
        vals = [f"{t['index']}", model_label(t), f"{len(t['placements'])}",
                _kg(total_weight_kg(t)), f"{rate * 100:.1f}%" if rate is not None else "—"]
        for j, v in enumerate(vals):
            _cell_text(tab.cell(i, j), v, size=9.5,
                       align=WD_ALIGN_PARAGRAPH.CENTER if j != 1 else None)
    last = 1 + len(trucks)
    total_w = sum(total_weight_kg(t) for t in trucks)
    total_pcs = sum(len(t["placements"]) for t in trucks)
    tot = [_bi("合计", "Tổng cộng"), "", f"{total_pcs}", _kg(total_w), ""]
    for j, v in enumerate(tot):
        _shade(tab.cell(last, j), "F5F5F5")
        _cell_text(tab.cell(last, j), v, size=9.5, bold=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER if j != 1 else None)
    for j, w in enumerate((1.6, 7.0, 1.8, 3.2, 3.3)):
        tab.columns[j].width = Cm(w)
    _para(doc, _bi("每辆车的摆放清单与 3D 装车示意（各一页）见后续各页。",
                   "Chi tiết xếp hàng và sơ đồ xếp hàng 3D của từng xe xem ở các trang sau (mỗi xe một trang)."),
          size=9, space_after=2)


def _add_truck_info(doc, t: dict, *, is_client: bool) -> None:
    tr = t["truck"]
    total = total_weight_kg(t)
    max_w = tr["maxWeight"]
    rate = loading_rate(t)
    weight_txt = _kg(total)
    if max_w > 0:
        pct = total / max_w * 100
        # 百分比必须写清是「占限重」，否则会和上一行「装载率」（体积口径）看串
        weight_txt += f" / {_kg(max_w)} kg（占限重 {pct:.1f}% / {pct:.1f}% tải trọng）"
        if total > max_w:
            weight_txt += _bi(" ⚠ 超重", " ⚠ Quá tải")
    info = [
        (_bi("车厢尺寸(长×宽×高 cm)", "Kích thước thùng xe (D×R×C cm)"),
         f"{_num(tr['L'])}×{_num(tr['W'])}×{_num(tr['H'])}"),
        (_bi("限重(kg)", "Tải trọng (kg)"), _kg(max_w) if max_w > 0 else "—"),
        (_bi("总重(kg)", "Tổng trọng lượng (kg)"), weight_txt),
        (_bi("装载率", "Tỷ lệ xếp hàng"), f"{rate * 100:.1f}%" if rate is not None else "—"),
        (_bi("件数", "Số kiện"), str(len(t["placements"]))),
    ]
    if not is_client:
        vol = truck_volume_cm3(t)
        used = used_volume_cm3(t)
        info.append((_bi("体积利用率(内部)", "Tỷ lệ thể tích (nội bộ)"),
                     f"{used / 1e6:.2f}/{vol / 1e6:.2f} m³" + (f"（{used / vol * 100:.1f}%）" if vol else "")))
    rows = [(_bi("车型", "Loại xe"), model_label(t))] + info
    tab = doc.add_table(rows=len(rows), cols=2)
    tab.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        _shade(tab.cell(i, 0), "F5F5F5")
        _cell_text(tab.cell(i, 0), k, size=9.5, bold=True)
        _cell_text(tab.cell(i, 1), v, size=9.5)
    tab.columns[0].width = Cm(5.6)
    tab.columns[1].width = Cm(11.4)
    _para(doc, "", space_after=2)


def _add_placements_table(doc, t: dict) -> None:
    _para(doc, _bi("摆放清单", "Danh sách xếp hàng"), size=11, bold=True, cn="黑体")
    heads = [_bi("序号", "STT"), _bi("货名", "Tên hàng"), _bi("尺寸(cm)", "Kích thước (cm)"),
             _bi("件重(kg)", "TL kiện (kg)"), _bi("位置(x,y,z cm)", "Vị trí (x,y,z cm)"),
             _bi("层", "Tầng"), _bi("不叠", "Không chồng")]
    tab = doc.add_table(rows=1 + len(t["placements"]), cols=len(heads))
    tab.style = "Table Grid"
    for j, h in enumerate(heads):
        _shade(tab.cell(0, j), "EDEDED")
        _cell_text(tab.cell(0, j), h, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    levels = layer_labels(t)
    for i, p in enumerate(t["placements"], start=1):
        vals = [
            f"{i}",
            p["name"],
            f"{_num(p['dx'])}×{_num(p['dz'])}×{_num(p['dy'])}",
            _kg(p["weight"]),
            f"{_num(p['x'])},{_num(p['y'])},{_num(p['z'])}",
            layer_text(t, p, levels),
            _bi("不叠", "Không chồng") if not p["stackable"] else "—",
        ]
        for j, v in enumerate(vals):
            _cell_text(tab.cell(i, j), v, size=9,
                       align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2, 3, 4, 5, 6) else None)
    for j, w in enumerate((1.1, 3.4, 3.3, 2.3, 3.1, 2.1, 1.7)):
        tab.columns[j].width = Cm(w)
    _para(doc, _AXIS_NOTE_BI, size=8.5, space_after=4)


def _add_unplaced_line(doc, t: dict) -> None:
    """未装配件一行（表下）：明确「需加车」，客户版也写（这是事实，不是内部口径）。"""
    items = t["unplaced"]
    if not items:
        return
    cn = "、".join(f"{u['name']}" + (f"（{u['reason']}）" if u["reason"] else "") for u in items)
    vi = "; ".join(f"{u['name']}" + (f" ({u['reason']})" if u["reason"] else "") for u in items)
    _para(doc, _bi(f"未能装配（需加车）：{cn}", f"Không xếp được (cần thêm xe): {vi}"),
          size=9.5, bold=True, space_after=4)


def _add_snapshot(doc, t: dict) -> None:
    """该车 3D 示意图：**只有真的拿到图才插图**（无图不留标题、不留空白页）。"""
    if not t["png"]:
        return
    _para(doc, _bi("装车示意图", "Sơ đồ xếp hàng"), size=11, bold=True, cn="黑体", space_after=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    try:
        p.add_run().add_picture(BytesIO(t["png"]), width=Cm(_IMG_WIDTH_CM))
    except Exception:  # noqa: BLE001 —— 图坏了也要能出单（与报价单盖章图同一纪律）
        _set_font(p.add_run(_bi("【示意图不可用】", "【Không hiển thị được sơ đồ】")), size=9)


def _add_truck_page(doc, t: dict, *, is_client: bool) -> None:
    _para(doc, _bi(f"第 {t['index']} 车", f"Xe số {t['index']}"), size=14, bold=True,
          cn="黑体", space_after=2)
    _para(doc, _truck_summary_line(t), size=9.5, space_after=6)
    _add_truck_info(doc, t, is_client=is_client)
    _add_placements_table(doc, t)
    _add_unplaced_line(doc, t)
    _add_snapshot(doc, t)


# ══════════════════════ 入口 ══════════════════════


def build(payload, profile: dict) -> bytes:
    """生成装箱单 docx（首页抬头 + 车辆汇总，随后**逐车一页**）。

    `payload` 可以是 pydantic 模型或等价的 dict。`trucks` 为空 → `ValueError`（端点折成 400）。
    """
    trucks = [_norm_truck(t, i) for i, t in enumerate(_g(payload, "trucks", []) or [], start=1)]
    if not trucks:
        raise ValueError("没有任何车辆数据（trucks 为空）")

    role = str(_g(payload, "role", "client") or "client").strip().lower()
    is_client = role != "internal"           # 只有显式 internal 才是内部版（其余一律按客户版出）
    job = _g(payload, "job", {}) or {}
    date_str = str(_g(job, "date", "") or "").strip() or datetime.date.today().strftime("%Y-%m-%d")

    doc = Document()
    _page_setup(doc)
    _add_letterhead(doc, profile)
    _add_title_block(doc, is_client=is_client, date_str=date_str)
    _add_job_info(doc, job, trucks)
    _add_summary_table(doc, trucks)
    for t in trucks:
        doc.add_page_break()                 # 逐车一页：第 1 个分页符隔开首页与第 1 车
        _add_truck_page(doc, t, is_client=is_client)
    _add_footer(doc, profile, f"{_TITLE_CN} {date_str}")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
